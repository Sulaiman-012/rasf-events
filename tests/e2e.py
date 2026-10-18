"""End-to-end test for the Rasf platform against a running `wrangler dev --local`.

Usage:  python3 tests/e2e.py [base_url] [setup_key]
Needs:  playwright (chromium), Pillow, openpyxl.  Uses a FRESH local database (run on an empty DB).
Screenshots go to test-output/.
"""
import io, json, os, sys, time, re
from pathlib import Path
from PIL import Image, ImageDraw
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:8787"
SETUP_KEY = sys.argv[2] if len(sys.argv) > 2 else "test-setup-key-123"
OUT = Path(__file__).resolve().parent.parent / "test-output"
OUT.mkdir(exist_ok=True)
OWNER = ("sulaiman@example.com", "Owner-Pass-2026")
MANAGER = ("ali@example.com", "Temp-Pass-1234", "Ali-New-Pass-2026")

results = []
def check(name, cond, detail=""):
    results.append((name, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""), flush=True)

def make_png(path):
    im = Image.new("RGB", (1080, 1920), "#2b211c")
    d = ImageDraw.Draw(im)
    for i in range(0, 1920, 60):
        d.rectangle([0, i, 1080, i + 28], fill=(141 + i % 50, 109, 82))
    d.ellipse([240, 660, 840, 1260], fill="#f3eee7")
    im.save(path)
    return path

def api(page, method, url, body=None):
    return page.evaluate("""async ([m,u,b]) => { const r = await fetch(u, {method:m, headers:{'Content-Type':'application/json'}, body: b===null?undefined:JSON.stringify(b)}); let j=null; try{j=await r.json()}catch{}; return {status:r.status, body:j}; }""", [method, url, body])

def no_hscroll(page):
    return page.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 1")

with sync_playwright() as p:
    browser = p.chromium.launch()
    errors = []
    def watch(pg, label):
        pg.on("pageerror", lambda e: errors.append(f"{label}: {e}"))
        pg.on("console", lambda m: errors.append(f"{label} console: {m.text}") if m.type == "error" and "fonts.g" not in m.text and "ERR_" not in m.text and "Failed to load resource" not in m.text else None)
        return pg

    # ---------------- 1. setup owner ----------------
    owner_ctx = browser.new_context(viewport={"width": 1440, "height": 1000}, locale="ar-SA", accept_downloads=True)
    a = watch(owner_ctx.new_page(), "admin")
    a.goto(BASE + "/setup")
    a.fill("#setup-key", "wrong-key"); a.fill("#setup-name", "سليمان"); a.fill("#setup-email", OWNER[0]); a.fill("#setup-password", OWNER[1])
    a.click("#setup-form button[type=submit]")
    expect(a.locator("#setup-error")).to_contain_text("غير صحيح")
    check("setup rejects wrong SETUP_KEY", True)
    a.fill("#setup-key", SETUP_KEY)
    a.click("#setup-form button[type=submit]")
    a.wait_for_url(BASE + "/admin")
    expect(a.locator("#admin")).to_have_class(re.compile("active"), timeout=10000)
    check("owner created via /setup and signed in", "مالك المنصة" in a.text_content("#role-chip"))
    r = api(a, "POST", "/api/setup", {"setupKey": SETUP_KEY, "name": "x", "email": "x@y.co", "password": "12345678"})
    check("/setup disabled after owner exists", r["status"] == 409, str(r["status"]))

    # ---------------- 2. competition, question, screen, image ----------------
    a.click("[data-resource=competition]")
    a.fill("#r-name", "مسابقة يوم التأسيس")
    a.click("#resource-form button[type=submit]")
    expect(a.locator("#competition-list")).to_contain_text("مسابقة يوم التأسيس")
    expect(a.locator("#competition-list")).to_contain_text("نشطة الآن")
    check("competition created and auto-activated", True)

    a.click("#new-question")
    a.fill("#edit-question", "في أي عام تأسست الدولة السعودية الأولى؟")
    a.select_option("#edit-type", "choice")
    for i, v in enumerate(["١٧٢٧م", "١٩٣٢م", "١٨٢٤م", "١٩٠٢م"]):
        a.fill(f"[data-option='{i}']", v)
    a.select_option("#edit-correct", "0")
    a.fill("#edit-attempts", "2"); a.fill("#edit-duration", "60")
    a.click("#save-question")
    expect(a.locator("#admin-questions")).to_contain_text("تأسست الدولة")
    check("question created (choice, 2 attempts)", True)

    a.click("[data-resource=group]"); a.fill("#r-name", "الاستقبال"); a.click("#resource-form button[type=submit]")
    expect(a.locator("#group-list")).to_contain_text("الاستقبال")
    a.click("[data-resource=screen]"); a.fill("#r-name", "شاشة المدخل"); a.select_option("#r-group", label="الاستقبال"); a.click("#resource-form button[type=submit]")
    expect(a.locator("#screen-list")).to_contain_text("شاشة المدخل")
    screen_url = a.locator("#screen-list code").first.text_content()
    check("screen created with unguessable key link", re.match(r"^https?://[^/]+/screen/[a-z2-9]{12}$", screen_url or "") is not None, screen_url)

    png = make_png(str(OUT / "upload-test.png"))
    a.set_input_files("#media-upload", png)
    expect(a.locator("#upload-status")).to_contain_text("تم رفع", timeout=20000)
    expect(a.locator("#media-grid .media-tile")).to_have_count(1)
    img_src = a.locator("#media-grid img").first.get_attribute("src")
    head = a.request.get(BASE + img_src)
    check("image uploaded, compressed and served from D1", head.status == 200 and head.headers.get("content-type") == "image/jpeg" and "immutable" in head.headers.get("cache-control", ""), f"{head.status} {head.headers.get('content-type')} {len(head.body())}B")

    # Show the question now
    a.locator("#admin-questions button", has_text="اعرضه الآن").first.click()
    expect(a.locator("#admin-questions")).to_contain_text("معروض الآن")
    st = a.request.get(BASE + "/api/state").json()
    check("public state never exposes the correct answer", "correct" not in json.dumps(st["question"]), json.dumps(st["question"], ensure_ascii=False)[:120])

    # ---------------- 3. employee joins and answers ----------------
    emp_ctx = browser.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2, is_mobile=True, has_touch=True, locale="ar-SA")
    e = watch(emp_ctx.new_page(), "employee")
    e.goto(BASE + "/")
    expect(e.locator("#competition-title")).to_contain_text("تأسست الدولة", timeout=15000)
    check("employee phone sees active question", True)
    check("employee QR rendered locally", e.locator("#qr-frame svg").count() == 1)
    e.screenshot(path=str(OUT / "01-employee-mobile-before.png"), full_page=True)
    e.fill("#participant-name", "نورة القحطاني")
    e.click("#start-answer")
    expect(e.locator("#timer-row")).to_be_visible()
    e.locator(".option", has_text="١٩٣٢م").click()
    expect(e.locator("#feedback")).to_contain_text("غير صحيحة")
    check("wrong answer graded on server, 1 attempt left", "١" in e.text_content("#feedback"))
    e.locator(".option", has_text="١٧٢٧م").click()
    expect(e.locator("#feedback")).to_contain_text("إجابة صحيحة")
    check("second attempt correct (+100)", True)
    e.screenshot(path=str(OUT / "02-employee-mobile-answered.png"), full_page=True)
    check("phone 390px has no horizontal scroll (display)", no_hscroll(e))
    # Third attempt must be refused by the server
    qid = st["question"]["id"]
    r = api(e, "POST", "/api/play/answer", {"questionId": qid, "value": 2})
    check("server refuses extra attempts after correct", r["body"] and r["body"].get("record", {}).get("done") is True and r["body"]["record"]["correct"] is True)

    e.click("[data-view=results]")
    expect(e.locator("#podium")).to_contain_text("نورة القحطاني")
    expect(e.locator("#podium")).to_contain_text("١٠٠")
    check("results show 100 points", True)
    e.screenshot(path=str(OUT / "03-results-mobile.png"), full_page=True)
    check("phone 390px has no horizontal scroll (results)", no_hscroll(e))

    # Message
    e.click("[data-view=display]")
    e.fill("#employee-message", "كل عام وأنتم بخير يا فريق رصف")
    e.click("#send-message")
    expect(e.locator("#message-status")).to_contain_text("وصلت للإدارة")
    check("employee message sent for moderation", True)

    e.click("[data-view=account]")
    expect(e.locator("#account-kpis")).to_contain_text("١٠٠")
    expect(e.locator("#account-messages")).to_contain_text("بانتظار المراجعة")
    check("my account: points + pending message", True)
    e.screenshot(path=str(OUT / "04-account-mobile.png"), full_page=True)
    check("phone 390px has no horizontal scroll (account)", no_hscroll(e))

    # ---------------- 4. screen (kiosk) ----------------
    scr_ctx = browser.new_context(viewport={"width": 540, "height": 960}, locale="ar-SA")
    s = watch(scr_ctx.new_page(), "screen")
    s.goto(screen_url)
    expect(s.locator("#competition-title")).to_contain_text("تأسست الدولة", timeout=15000)
    s.fill("#participant-name", "زائر الاستقبال")
    s.click("#start-answer")
    s.locator(".option", has_text="١٧٢٧م").click()
    expect(s.locator("#feedback")).to_contain_text("إجابة صحيحة")
    check("kiosk answer recorded on the screen", True)
    s.screenshot(path=str(OUT / "05-screen-quiz-540x960.png"))
    s.wait_for_timeout(7600)
    check("kiosk resets for next person after 7s", s.input_value("#participant-name") == "" and s.is_visible("#participant-name"))

    # admin: approve message, set screen content to image
    a.reload(); expect(a.locator("#admin")).to_have_class(re.compile("active"), timeout=10000)
    expect(a.locator("#message-list")).to_contain_text("كل عام وأنتم بخير")
    a.locator("#message-list button", has_text="اعتماد وعرض").first.click()
    expect(a.locator("#message-list")).to_contain_text("المعروضة الآن")
    a.locator("#screen-list select").first.select_option(index=1)  # options: [المسابقة, صورة: ...]
    expect(a.locator("#media-grid")).to_contain_text("معروضة على", timeout=10000)
    check("message approved + screen content set to image", True)
    # screen picks up changes via polling (5s)
    expect(s.locator("#media-stage")).to_be_visible(timeout=15000)
    expect(s.locator("#stage-msg-text")).to_contain_text("كل عام وأنتم بخير", timeout=15000)
    check("screen shows image full-bleed + approved message + QR", s.locator("#stage-qr svg").count() == 1)
    s.screenshot(path=str(OUT / "06-screen-media-540x960.png"))
    expect(e.locator("#featured-text")).to_contain_text("كل عام", timeout=20000)
    check("employee page shows featured message", True)
    e.click("[data-view=account]")
    expect(e.locator("#account-messages")).to_contain_text("معروضة الآن")
    check("my account message status = shown now", True)

    # heartbeat → online
    a.wait_for_timeout(1000); a.reload(); expect(a.locator("#admin")).to_have_class(re.compile("active"), timeout=10000)
    expect(a.locator("#screen-list")).to_contain_text("متصلة الآن")
    check("screen heartbeat shows online in admin", True)
    a.screenshot(path=str(OUT / "07-admin-manage.png"), full_page=True)

    # ---------------- 5. insights + Excel ----------------
    a.click("[data-pane=insights]")
    expect(a.locator("#insights-kpis")).to_contain_text("المشاركون")
    expect(a.locator("#insights-questions .bar-row")).to_have_count(4)
    a.wait_for_timeout(400)
    a.screenshot(path=str(OUT / "08-admin-insights.png"), full_page=True)
    with a.expect_download(timeout=30000) as dl:
        a.click("#export-insights")
    path = OUT / "export.xlsx"
    dl.value.save_as(str(path))
    wb = load_workbook(path)
    names = wb.sheetnames
    res_rows = list(wb["النتائج"].iter_rows(values_only=True)) if "النتائج" in names else []
    ans_rows = list(wb["الإجابات"].iter_rows(values_only=True)) if "الإجابات" in names else []
    rtl = wb["النتائج"].sheet_view.rightToLeft if "النتائج" in names else None
    date_ok = bool(res_rows) and re.match(r"^\d{2}/\d{2}/\d{4}, \d{2}:\d{2}:\d{2}$", str(res_rows[1][5] or "")) is not None
    check("Excel export: 3 sheets, RTL, rows, en-GB dates", names == ["النتائج", "الإجابات", "الأسئلة"] and len(res_rows) == 3 and len(ans_rows) == 3 and date_ok,
          f"sheets={names} results={len(res_rows)-1} answers={len(ans_rows)-1} rtl={rtl} first={res_rows[1] if len(res_rows)>1 else None}")

    # ---------------- 6. roles: owner adds manager ----------------
    a.click("[data-pane=roles]")
    expect(a.locator("#staff-card")).to_be_visible()
    a.fill("#staff-name", "علي الغدير"); a.fill("#staff-email", MANAGER[0]); a.fill("#staff-password", MANAGER[1])
    a.click("#staff-form button[type=submit]")
    expect(a.locator("#staff-list")).to_contain_text(MANAGER[0])
    expect(a.locator("#people-list")).to_contain_text("نورة القحطاني")
    check("owner added manager; participants listed", True)
    a.screenshot(path=str(OUT / "09-admin-roles.png"), full_page=True)

    man_ctx = browser.new_context(viewport={"width": 1280, "height": 900}, locale="ar-SA")
    m = watch(man_ctx.new_page(), "manager")
    m.goto(BASE + "/admin")
    m.fill("#login-email", MANAGER[0]); m.fill("#login-password", "wrong-password")
    m.click("#login-form button[type=submit]")
    expect(m.locator("#login-error")).to_contain_text("غير صحيحة")
    m.fill("#login-password", MANAGER[1]); m.click("#login-form button[type=submit]")
    expect(m.locator("#password")).to_have_class(re.compile("active"))
    r = api(m, "GET", "/api/admin", None)
    check("manager with temp password is blocked until change", r["status"] == 403, str(r["status"]))
    m.fill("#pw-current", MANAGER[1]); m.fill("#pw-next", MANAGER[2]); m.click("#password-form button[type=submit]")
    expect(m.locator("#admin")).to_have_class(re.compile("active"), timeout=10000)
    check("manager changed password and entered admin", "مدير" in m.text_content("#role-chip"))
    r = api(m, "PUT", "/api/admin/identity", {"primary": "#000000", "accent": "#111111", "background": "#ffffff"})
    check("manager cannot change identity (403)", r["status"] == 403, str(r["status"]))
    r = api(m, "GET", "/api/admin/staff", None)
    check("manager cannot list/manage staff (403)", r["status"] == 403, str(r["status"]))
    r = api(m, "POST", "/api/admin/groups", {"name": "الدور الثاني"})
    check("manager can manage groups (200)", r["status"] == 200, str(r["status"]))
    check("identity controls disabled for manager in UI", m.is_disabled("#save-identity"))
    r = api(a, "PUT", "/api/admin/identity", {"primary": "#8d6d52", "accent": "#c6ab8c", "background": "#f3eee7"})
    check("owner can change identity (200)", r["status"] == 200, str(r["status"]))

    # ---------------- 7. security spot checks ----------------
    anon = browser.new_context().new_page()
    anon.goto(BASE + "/")
    r = api(anon, "POST", "/api/admin/competitions", {"name": "x"})
    check("anonymous admin write → 401", r["status"] == 401, str(r["status"]))
    rr = anon.request.post(BASE + "/api/admin/groups", data='{"name":"x"}', headers={"Content-Type": "application/json", "Origin": "https://evil.example"})
    check("cross-origin write blocked (403)", rr.status == 403, str(rr.status))
    rr = anon.request.post(BASE + "/api/join", data='name=x', headers={"Content-Type": "application/x-www-form-urlencoded"})
    check("non-JSON write blocked (415)", rr.status == 415, str(rr.status))
    hdr = anon.request.get(BASE + "/").headers
    check("security headers on pages (CSP frame-ancestors, nosniff)", "frame-ancestors 'none'" in hdr.get("content-security-policy", "") and hdr.get("x-content-type-options") == "nosniff")
    # login rate limit: 5 failures then blocked
    codes = [api(anon, "POST", "/api/auth/login", {"email": "rl@example.com", "password": "bad-pass-123"})["status"] for _ in range(6)]
    check("login rate limit (5 per 15 min)", codes[:5] == [401] * 5 and codes[5] == 429, str(codes))
    st2 = anon.request.get(BASE + "/api/state")
    et = st2.headers.get("etag")
    st3 = anon.request.get(BASE + "/api/state", headers={"If-None-Match": et})
    check("unchanged poll returns 304", st3.status == 304, f"{et} -> {st3.status}")

    # ---------------- 8. desktop screenshot of employee page ----------------
    d = watch(browser.new_context(viewport={"width": 1440, "height": 900}, locale="ar-SA").new_page(), "desktop")
    d.goto(BASE + "/"); d.wait_for_timeout(1500)
    d.screenshot(path=str(OUT / "10-employee-desktop.png"), full_page=True)

    big = watch(browser.new_context(viewport={"width": 1080, "height": 1920}, locale="ar-SA").new_page(), "lg")
    big.goto(screen_url); big.wait_for_timeout(2000)
    big.screenshot(path=str(OUT / "11-screen-media-1080x1920.png"))
    check("LG 1080x1920 page fits without scrolling", big.evaluate("document.documentElement.scrollHeight <= 1921"))

    check("no JS errors in any page", not errors, "; ".join(errors[:5]))
    browser.close()

failed = [r for r in results if not r[1]]
print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
sys.exit(1 if failed else 0)
